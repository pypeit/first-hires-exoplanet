""" Two slides on the project for a general audience (phase 3, prompt 16).

Writes ``docs/slides/public_summary.pptx`` (16:9):

    1. "You could do most of this yourself": the ten prompts that steered the
       project, and how few decisions it took.
    2. "A teacher could lead a class through it": the project as six lessons.

The counts on slide 1 (steps, decisions, days) are taken from the prompt docs
in ``claude_prompts/`` at build time, not typed:

    steps      numbered items under each doc's ``## Prompts``
    decisions  answers recorded in the Q&A sections (``>A.`` or ``Answer:``)
    days       first to last dated log entry

The prompts quoted on slide 1 are shortened from the originals in those
docs; ``PROMPTS`` below records where each one is.

Layout and the 20-pt floor come from ``make_slides.py``.

Run with:

    conda run -n pypeit14 python -m first_hires_exoplanet.make_public_slides

"""

# Standard imports
import os
import re
import glob
import argparse
import datetime

from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from first_hires_exoplanet.make_slides import (Deck, rgb, _pt, MIN_PT, TAKE_TOP, ROOT,
                                               BLUE, ORANGE, AQUA, INK, INK2, W)


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

PROMPT_DIR = os.path.join(ROOT, 'claude_prompts')
OUT = os.path.join(ROOT, 'docs', 'slides', 'public_summary.pptx')

#: The ten prompts that steered the project: (shortened text, where it is)
PROMPTS = [
    ('Find the first exoplanet discovered with Keck/HIRES.', 'context_prompts.md #1'),
    ('Ask me another round of questions.', 'context_prompts.md #2-4'),
    ('Make clear it is the star that moves; the planet is invisible.', 'context_prompts.md #6'),
    ('Write a plan for working on the 1998 data.', 'context_prompts.md #7'),
    ('Check that PypeIt is ready for the original HIRES detector.', 'data_phase1_prompt.md #1'),
    ('Fix the issues you have identified.', 'data_phase1_prompt.md #6'),
    ('Take account of pyodine: adapt it rather than build it.', 'context_prompts.md #8'),
    ('Reread the file and execute prompt #N.', 'every step'),
    ('Answer: (a)', 'every decision'),
    ('Draft the report to the PypeIt team; show it to me first.', 'data_phase3_prompt.md #10'),
]

#: The project as six lessons: (title, what the students do)
LESSONS = [
    ('The wobble', 'A planet tugs its star; we see the star move'),
    ('The archive', 'Find the 1998 Keck images: they are public'),
    ('Light into spectra', 'Turn raw images into a barcoded rainbow'),
    ('Lamps fall short', 'Our first try missed the planet. Why?'),
    ('A ruler of iodine', 'Gas in the light path marks every exposure'),
    ('Find the planet', 'Fit a 3.1-day wave; compare with 1998'),
]


# ---------------------------------------------------------------------------
# Counts, from the prompt docs
# ---------------------------------------------------------------------------

def _section(text, name):
    """ The body of a ``## name`` section, up to the next ``## `` heading. """
    m = re.search(r'^## {:s}\s*$(.*?)(?=^## |\Z)'.format(re.escape(name)), text, re.M | re.S)
    return m.group(1) if m else ''


def counts(prompt_dir=PROMPT_DIR):
    """ Steps, decisions and elapsed days, from every prompt doc. """
    steps, decisions, dates = 0, 0, []
    for f in sorted(glob.glob(os.path.join(prompt_dir, '*.md'))):
        with open(f) as fh:
            text = fh.read()
        steps += len(re.findall(r'^\s*\d+\.\s', _section(text, 'Prompts'), re.M))
        # Every "## Q&A..." section (context_prompts.md has four rounds)
        for qa in re.findall(r'^## Q&A.*?$(.*?)(?=^## (?!Q&A)|\Z)', text, re.M | re.S):
            decisions += len(re.findall(r'^\s*(?:>\s*A\.|Answer:)', qa, re.M))
        logs = _section(text, 'Logs')
        dates += re.findall(r'^### (\d{4}-\d{2}-\d{2})', logs, re.M)
    d = sorted(datetime.date.fromisoformat(x) for x in dates)
    return dict(steps=steps, decisions=decisions, first=d[0], last=d[-1],
                days=(d[-1] - d[0]).days + 1)


# ---------------------------------------------------------------------------
# The slides
# ---------------------------------------------------------------------------

def _box(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    return shp


def _fill(shp, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    """ Write (text, size, bold, color) paragraphs into a shape. """
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.15)
    for i, (txt, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = txt
        r.font.size, r.font.bold = _pt(size), bold
        r.font.color.rgb = rgb(color)


def slide_prompts(d, C):
    s = d.prs.slides.add_slide(d.blank)
    d._chrome(s, 'You could do most of this yourself')

    # Left: the ten prompts
    d._text(s, Inches(0.55), Inches(1.15), Inches(8.6), Inches(0.5),
            [('Ten prompts steered it (shortened):', dict(color=INK2))])
    paras = []
    for i, (txt, _) in enumerate(PROMPTS, start=1):
        star = i in (8, 9)
        paras.append(('{:d}.  {:s}'.format(i, txt),
                      dict(bold=star, color=BLUE if star else INK, space=3)))
    d._text(s, Inches(0.55), Inches(1.6), Inches(8.7), Inches(4.7), paras)

    # Right: the numbers
    tiles = [
        ('{:d} days'.format(C['days']), '{:s} – {:s}'.format(
            C['first'].strftime('%b %-d'), C['last'].strftime('%b %-d, %Y'))),
        ('{:d} steps'.format(C['steps']), '“execute prompt #N”'),
        ('{:d} decisions'.format(C['decisions']), 'often just “(a)”'),
        ('All free', 'public data and software'),
    ]
    x, y0, tw, th, gap = Inches(9.45), Inches(1.2), Inches(3.48), Inches(1.15), Inches(0.12)
    for k, (big, small) in enumerate(tiles):
        _fill(_box(s, x, y0 + k * (th + gap), tw, th, 'eef4fc'),
              [(big, 28, True, BLUE), (small, MIN_PT, False, INK2)])

    d._takeaway(s, 'I set the goal and made the calls. The AI wrote the plans, the code and the '
                   'reports — and re-found the planet at the published 72 m/s.')
    return s


def slide_lessons(d):
    s = d.prs.slides.add_slide(d.blank)
    d._chrome(s, 'A teacher could lead a class through it')
    colors = [AQUA, AQUA, '8a8a86', '8a8a86', BLUE, BLUE]
    bw, bh, gx, gy = Inches(4.0), Inches(1.75), Inches(0.26), Inches(0.2)
    x0, y0 = Inches(0.4), Inches(1.25)
    for k, ((title, body), col) in enumerate(zip(LESSONS, colors)):
        r, c = divmod(k, 3)
        x, y = x0 + c * (bw + gx), y0 + r * (bh + gy)
        _box(s, x, y, Inches(0.12), bh, col)
        _fill(_box(s, x + Inches(0.12), y, bw - Inches(0.12), bh, 'f3f3f1'),
              [('Lesson {:d} · {:s}'.format(k + 1, title), 22, True, INK),
               (body, MIN_PT, False, INK2)])
    d._text(s, Inches(0.55), y0 + 2 * bh + gy + Inches(0.12), Inches(12.2), Inches(1.0), [
        ('What I would hand the teacher: this repository (every plan, question and answer), the ten prompts, '
         'the public write-up for students to read, and the judgment calls: when a result looks wrong, and why.',
         dict(size=MIN_PT)),
    ])
    d._takeaway(s, 'The scarce skill was knowing what to ask and when to doubt an answer. '
                   'That can be taught.')
    return s


def build(out):
    C = counts()
    print('  counts: {:d} steps, {:d} decisions, {:d} days ({:s} to {:s})'.format(
        C['steps'], C['decisions'], C['days'], str(C['first']), str(C['last'])))
    d = Deck()
    slide_prompts(d, C)
    slide_lessons(d)
    d.save(out)


def main():
    p = argparse.ArgumentParser(description='Build the two public slides (phase 3, prompt 16)')
    p.add_argument('--out', type=str, default=OUT)
    build(p.parse_args().out)


if __name__ == '__main__':
    main()
